from typing import Iterable, Tuple

import openpyxl
from tqdm import tqdm

from .. import logger
from .Document import Document


class XlsDocument(Document):
    def iterate_raw_text(self) -> Iterable[Tuple[str, dict]]:
        try:
            wb = openpyxl.load_workbook(self.get_abs_path(), read_only=True, data_only=True)
        except Exception:
            logger.warning("Error while reading the file. Skipping")
            return None, {"ocr_used": False}

        logger.info(f"Reading {len(wb.worksheets)} pages excel file")
        all_text = []
        for sheet in tqdm(wb.worksheets):
            for row in sheet.iter_rows(values_only=True):
                row_text = [str(cell) for cell in row if cell is not None]
                if row_text:
                    all_text.append(" ".join(row_text))
            yield "\n".join(all_text).strip(), {"ocr_used": False}
